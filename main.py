# Внешние библиотеки
# Документы по шаблону (template.docx)
# Excel (openpyxl)
#pip install openpyxl
#pip freeze > requirements.txt

#Пустой Excel -файл
'''from openpyxl import Workbook

wb = Workbook()

ws = wb.active
ws.title = 'Отчет'

wb.save('Docs/report.xlsx')'''


from itertools import count

#Запись данных в существующий файл
from openpyxl import load_workbook

#открываем(загружаем) рабочую книгу
wb = load_workbook('Docs/report.xlsx')

#активный лист
ws = wb.active

#способы записи
ws['F1'] = 'Привет мир'    #в какую ячейку писать (латинские буквы)
ws.cell(row=1, column=3, value='Hello!')

wb.save('Docs/newtable.xlsx')